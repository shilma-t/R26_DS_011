"""
merge_rediarized_worksheet.py
Rebuild the verification worksheet after re-diarizing a subset of calls,
WITHOUT losing verification labels for the calls that did not change.

Rows for unchanged calls are carried over exactly as the human entered them.
Rows for re-diarized calls are regenerated blank, because their speaker
labels have been reassigned and the old answers no longer refer to the
same voices.

Writes to a NEW file. The existing worksheet is never modified.

Usage:
  python scripts/pipeline/merge_rediarized_worksheet.py --prefix 890001
"""

import os
import argparse

import pandas as pd

_ROOT   = r"C:\Users\shafr\Desktop\Uni\RESEARCH\R26_DS_011"
REPORTS = os.path.join(_ROOT, "reports", "pp2")

MANIFEST  = os.path.join(REPORTS, "dmc_diarization_manifest.csv")
WORKSHEET = os.path.join(REPORTS, "dmc_verify_worksheet.xlsx")
OUTPUT    = os.path.join(REPORTS, "dmc_verify_worksheet_v2.xlsx")

N_SAMPLES   = 3
MIN_SEG_SEC = 1.0


def safe_name(text):
    return "".join(ch if (ch.isalnum() or ch in "._-") else "_" for ch in str(text))


def mmss(seconds):
    seconds = float(seconds)
    minutes = int(seconds // 60)
    return f"{minutes}:{seconds - minutes * 60:04.1f}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--prefix", required=True,
                        help="Recording-id prefix of the re-diarized calls")
    args = parser.parse_args()

    for path in (MANIFEST, WORKSHEET):
        if not os.path.exists(path):
            raise SystemExit(f"Not found: {path}")

    # --- existing human work -------------------------------------------------
    old = pd.read_excel(WORKSHEET, dtype=str)
    old = old[old["speaker"].notna()
              & old["speaker"].astype(str).str.startswith("SPEAKER")].copy()

    changed = old["recording_id"].astype(str).str.startswith(args.prefix)
    keep = old[~changed].copy()
    drop = old[changed]

    print(f"Existing worksheet : {len(old)} rows, {old['recording_id'].nunique()} calls")
    print(f"  carried over     : {len(keep)} rows "
          f"({keep['recording_id'].nunique()} calls, labels preserved)")
    print(f"  to be regenerated: {len(drop)} rows "
          f"({drop['recording_id'].nunique()} calls, labels cleared)")

    if keep["is_caller"].isna().any():
        blanks = keep["is_caller"].isna().sum()
        print(f"\n  WARNING: {blanks} carried-over row(s) have no label.")

    # Reuse each re-diarized call's original clip number so filenames stay stable
    clip_index = {}
    for row in drop.itertuples():
        prefix = str(row.clip_file).split("_", 1)[0]
        if prefix.isdigit():
            clip_index[str(row.recording_id)] = int(prefix)

    # --- fresh rows from the new manifest ------------------------------------
    man = pd.read_csv(MANIFEST, dtype={"recording_id": str, "filename": str})
    man = man[man["speaker"].notna() & (man["speaker"] != "ERROR")]
    for col in ("start", "end", "duration"):
        man[col] = pd.to_numeric(man[col], errors="coerce")
    man = man.dropna(subset=["start", "end", "duration"])
    man = man[man["recording_id"].astype(str).str.startswith(args.prefix)]

    fresh = []
    for rec_id, call in man.groupby("recording_id", sort=True):
        call = call.sort_values("start")
        index = clip_index.get(str(rec_id))
        if index is None:
            raise SystemExit(
                f"No previous clip number found for {rec_id}. "
                "Cannot keep clip filenames stable - aborting rather than guessing."
            )

        opener = call[call["duration"] >= MIN_SEG_SEC]
        first_speaker = (opener["speaker"].iloc[0]
                         if len(opener) else call["speaker"].iloc[0])

        totals = call.groupby("speaker")["duration"].sum().sort_values(ascending=False)
        grand = totals.sum()

        for speaker in totals.index:
            sgrp = call[call["speaker"] == speaker]
            usable = sgrp[sgrp["duration"] >= MIN_SEG_SEC]
            if usable.empty:
                usable = sgrp
            picks = usable.nlargest(N_SAMPLES, "duration").sort_values("start")

            secs = totals[speaker]
            fresh.append({
                "recording_id": rec_id,
                "clip_file":    f"{index:03d}_{safe_name(rec_id)}__{speaker}.wav",
                "speaker":      speaker,
                "play_these":   "   ".join(f"{mmss(r.start)}-{mmss(r.end)}"
                                           for r in picks.itertuples()),
                "total_speech": f"{secs:.1f}s ({100 * secs / grand:.0f}%)",
                "speaks_first": "YES" if speaker == first_speaker else "",
                "is_caller":    "",
            })

    fresh_df = pd.DataFrame(fresh)
    print(f"\nRegenerated        : {len(fresh_df)} rows "
          f"({fresh_df['recording_id'].nunique()} calls) - blank, ready to relabel")

    merged = pd.concat([keep, fresh_df], ignore_index=True)

    # --- write ---------------------------------------------------------------
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verify"

    headers = ["recording_id", "clip_file", "speaker", "play_these",
               "total_speech", "speaks_first", "is_caller"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="12403C")
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    thin   = Side(style="thin", color="7A8987")
    todo   = PatternFill("solid", fgColor="FFD9D2")   # needs relabelling
    done   = PatternFill("solid", fgColor="E4EFE7")   # already verified
    band   = PatternFill("solid", fgColor="EDF2F1")

    previous_id = None
    banded = False
    for row in merged.to_dict("records"):
        ws.append([row.get(h, "") for h in headers])
        r = ws.max_row

        if row["recording_id"] != previous_id:
            banded = not banded
            previous_id = row["recording_id"]
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).border = Border(top=thin)

        needs_work = not str(row.get("is_caller") or "").strip()

        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name="Arial", size=10)
            if banded and c != 7:
                cell.fill = band

        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2).font = Font(name="Consolas", size=9)
        ws.cell(r, 4).font = Font(name="Consolas", size=10)
        ws.cell(r, 6).font = Font(name="Arial", size=10, bold=True, color="A8600F")
        ws.cell(r, 7).fill = todo if needs_work else done

    widths = {"A": 32, "B": 44, "C": 13, "D": 34, "E": 15, "F": 13, "G": 11}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    note_row = ws.max_row + 3
    notes = [
        "WHAT CHANGED",
        f"The {fresh_df['recording_id'].nunique()} calls starting {args.prefix} were resampled from 8 kHz to",
        "16 kHz and re-diarized, so their speaker labels were reassigned.",
        "Those rows are PINK and blank - relabel them from the new clips.",
        "",
        "GREEN rows are your existing verified answers, carried over untouched.",
        "",
        "Values: yes / no / yes-mixed / no-mixed / noise / unclear",
    ]
    for i, line in enumerate(notes):
        ws.cell(note_row + i, 1, line).font = Font(name="Arial", size=10, bold=(i == 0))

    wb.save(OUTPUT)

    print(f"\nWrote {OUTPUT}")
    print(f"  {len(merged)} rows total")
    print(f"  {(merged['is_caller'].fillna('').str.strip() != '').sum()} already labelled (green)")
    print(f"  {(merged['is_caller'].fillna('').str.strip() == '').sum()} to relabel (pink)")
    print("\nYour original worksheet was NOT modified.")


if __name__ == "__main__":
    main()
