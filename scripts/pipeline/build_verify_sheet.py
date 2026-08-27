"""
build_verify_sheet.py
Rebuild the diarization verification worksheet WITH timestamps.

One row per speaker (not per call), so each row is a single question:
"play these times - is this the caller?"

The timestamps given are that speaker's longest, cleanest segments, which
are the easiest to identify a voice from.

Usage:
  python scripts/pipeline/build_verify_sheet.py            # pilot
  python scripts/pipeline/build_verify_sheet.py --full     # full manifest

Output:
  reports/pp2/dmc_pilot_verify.xlsx
"""

import os
import argparse

import pandas as pd

_ROOT   = r"C:\Users\shafr\Desktop\Uni\RESEARCH\R26_DS_011"
REPORTS = os.path.join(_ROOT, "reports", "pp2")

PILOT_CSV = os.path.join(REPORTS, "dmc_pilot_diarization.csv")
FULL_CSV  = os.path.join(REPORTS, "dmc_diarization_manifest.csv")

N_SAMPLES   = 3     # how many timestamp ranges to offer per speaker
MIN_SEG_SEC = 1.0   # ignore backchannels / overlap fragments


def safe_name(text):
    """Filename-safe recording id. Must match make_verify_clips.py."""
    return "".join(ch if (ch.isalnum() or ch in "._-") else "_" for ch in str(text))


def mmss(seconds):
    """Format seconds as m:ss so it matches a media player's display."""
    seconds = float(seconds)
    minutes = int(seconds // 60)
    rest = seconds - minutes * 60
    return f"{minutes}:{rest:04.1f}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--full", action="store_true")
    parser.add_argument("--force", action="store_true",
                        help="Overwrite a worksheet that already has verified rows")
    args = parser.parse_args()

    src = FULL_CSV if args.full else PILOT_CSV
    if not os.path.exists(src):
        raise SystemExit(f"Manifest not found: {src}")

    df = pd.read_csv(src, dtype={"recording_id": str, "filename": str})
    df = df[df["speaker"].notna() & (df["speaker"] != "ERROR")]
    for col in ("start", "end", "duration"):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df.dropna(subset=["start", "end", "duration"])

    rows = []

    # sort=False and start=1 must match make_verify_clips.py so clip
    # numbering lines up with the rows in this sheet
    for index, (rec_id, call) in enumerate(df.groupby("recording_id", sort=False), start=1):
        fname = call["filename"].iloc[0]
        call = call.sort_values("start")

        # Who opens the call (first segment of real length) -> operator hypothesis
        opener = call[call["duration"] >= MIN_SEG_SEC]
        first_speaker = (opener["speaker"].iloc[0]
                         if len(opener) else call["speaker"].iloc[0])

        totals = call.groupby("speaker")["duration"].sum().sort_values(ascending=False)
        grand_total = totals.sum()

        for speaker in totals.index:
            sgrp = call[call["speaker"] == speaker]

            # Longest segments are the easiest to identify a voice from
            usable = sgrp[sgrp["duration"] >= MIN_SEG_SEC]
            if usable.empty:
                usable = sgrp
            picks = usable.nlargest(N_SAMPLES, "duration").sort_values("start")

            times = "   ".join(
                f"{mmss(r.start)}-{mmss(r.end)}" for r in picks.itertuples()
            )

            secs = totals[speaker]
            rows.append({
                "recording_id":  rec_id,
                "clip_file":     f"{index:03d}_{safe_name(rec_id)}__{speaker}.wav",
                "speaker":       speaker,
                "play_these":    times,
                "total_speech":  f"{secs:.1f}s ({100 * secs / grand_total:.0f}%)",
                "speaks_first":  "YES" if speaker == first_speaker else "",
                "is_caller":     "",
            })

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verify"

    headers = ["recording_id", "clip_file", "speaker", "play_these",
               "total_speech", "speaks_first", "is_caller"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="12403C")
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    thick = Side(style="thin", color="7A8987")
    fill_in = PatternFill("solid", fgColor="FFF3C4")
    band    = PatternFill("solid", fgColor="EDF2F1")

    previous_id = None
    banded = False
    for row in rows:
        ws.append([row[h] for h in headers])
        r = ws.max_row

        # Alternate shading per CALL so speaker groups read as blocks
        if row["recording_id"] != previous_id:
            banded = not banded
            previous_id = row["recording_id"]
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).border = Border(top=thick)

        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name="Arial", size=10)
            if banded and c != 7:
                cell.fill = band

        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2).font = Font(name="Consolas", size=9)
        ws.cell(r, 4).font = Font(name="Consolas", size=10)
        ws.cell(r, 6).font = Font(name="Arial", size=10, bold=True, color="A8600F")
        ws.cell(r, 7).fill = fill_in

    widths = {"A": 32, "B": 44, "C": 13, "D": 34, "E": 15, "F": 13, "G": 11}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # Instruction block below the table
    note_row = ws.max_row + 3
    notes = [
        "HOW TO FILL THIS IN",
        "1. Open audio/dmc_verify/ and play the file named in 'clip_file'.",
        "   Each clip is ~12s of ONLY that speaker. Clips are numbered in",
        "   this sheet's order, so you can work straight down the list.",
        "2. Operator = formal, asks short questions.",
        "   Caller = explains the situation, gives location, sounds stressed.",
        "3. Write yes or no in the yellow 'is_caller' column.",
        "4. Normally ONE speaker per recording_id is yes. Mark a second row",
        "   yes only if diarization split one voice across two labels.",
        "",
        "'play_these' gives the same moments as timestamps in the original",
        "recording, if you ever need to hear them in context.",
        "'speaks_first' marks whoever opens the call - verified at only 60%",
        "on the pilot, so treat it as a hint, not an answer.",
    ]
    for i, line in enumerate(notes):
        cell = ws.cell(note_row + i, 1, line)
        cell.font = Font(name="Arial", size=10, bold=(i == 0))

    name = "dmc_verify_worksheet.xlsx" if args.full else "dmc_pilot_verify.xlsx"
    out = os.path.join(REPORTS, name)

    # Never silently destroy human verification work
    if os.path.exists(out) and not args.force:
        try:
            existing = pd.read_excel(out)
            done = existing["is_caller"].notna().sum() if "is_caller" in existing else 0
        except Exception:
            done = 0
        if done:
            raise SystemExit(
                f"REFUSING TO OVERWRITE: {out}\n"
                f"It already has {done} verified row(s) filled in.\n"
                "Rename or move that file first, or pass --force to discard it."
            )

    wb.save(out)

    calls = len({r["recording_id"] for r in rows})
    print(f"Wrote {out}")
    print(f"{calls} calls, {len(rows)} speaker rows to check.")
    print("Clips referenced in column B live in audio/dmc_verify/")


if __name__ == "__main__":
    main()
